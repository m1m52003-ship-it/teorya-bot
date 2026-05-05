from docx import Document
import os
import re

# =========================
# إعدادات الملفات
# =========================
QUESTIONS_FILE = "questions.docx"
ANSWERS_FILE = "answers.docx"
IMAGES_FOLDER = "images"


# =========================
# خريطة الإجابات
# الإجابة الصحيحة تعتمد على الاختيار:
# أ / ب / ج / د
# =========================
ANSWER_MAP = {
    "أ": 0,
    "ا": 0,
    "إ": 0,
    "آ": 0,
    "ب": 1,
    "ج": 2,
    "د": 3,
}


# =========================
# خريطة تحويل الاختيارات الرقمية إلى أحرف
# 1 => أ
# 2 => ب
# 3 => ج
# 4 => د
# =========================
NUMBER_TO_LETTER = {
    "1": "أ",
    "2": "ب",
    "3": "ج",
    "4": "د",
}


# =========================
# تنظيف النص
# =========================
def clean_text(text):
    if not text:
        return ""

    text = text.replace("\t", " ")
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


# =========================
# تنظيف حرف الإجابة
# =========================
def normalize_answer(ans):
    ans = clean_text(ans)

    ans = ans.replace(".", "")
    ans = ans.replace("-", "")
    ans = ans.replace(":", "")
    ans = ans.replace(")", "")
    ans = ans.replace("(", "")
    ans = ans.strip()

    if ans in ["ا", "إ", "آ"]:
        ans = "أ"

    return ans


# =========================
# قراءة الإجابات من answers.docx
# الشكل المتوقع:
# 1.1 ج
# 1.2 أ
# =========================
def read_answers(file_path):
    answers = {}

    doc = Document(file_path)

    # قراءة الفقرات العادية
    for paragraph in doc.paragraphs:
        line = clean_text(paragraph.text)

        if not line:
            continue

        match = re.search(r"(\d+\.\d+)\s*([أاإآبجد])", line)

        if match:
            q_id = match.group(1)
            ans = normalize_answer(match.group(2))

            if ans in ANSWER_MAP:
                answers[q_id] = {
                    "letter": ans,
                    "index": ANSWER_MAP[ans]
                }

    # قراءة الجداول
    for table in doc.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]

            if len(cells) < 2:
                continue

            q_id = None
            ans = None

            for cell in cells:
                id_match = re.search(r"\d+\.\d+", cell)
                if id_match:
                    q_id = id_match.group(0)

                answer_match = re.search(r"[أاإآبجد]", cell)
                if answer_match:
                    possible_ans = normalize_answer(answer_match.group(0))
                    if possible_ans in ANSWER_MAP:
                        ans = possible_ans

            if q_id and ans:
                answers[q_id] = {
                    "letter": ans,
                    "index": ANSWER_MAP[ans]
                }

    return answers


# =========================
# هل السطر بداية سؤال؟
# مثال:
# 5.12 نص السؤال
# =========================
def is_question_line(text):
    return re.match(r"^\d+\.\d+\s+", text) is not None


# =========================
# هل السطر اختيار؟
# يقبل:
# أ. ...
# ب . ...
# ج) ...
# د - ...
#
# وأيضًا:
# 1. ...
# 2. ...
# 3. ...
# 4. ...
# =========================
def is_option_line(text):
    text = clean_text(text)

    arabic_option = re.match(r"^[أاإآبجد]\s*[\.\-\)]?\s*", text)
    numeric_option = re.match(r"^[1-4]\s*[\.\-\)]\s*", text)

    return arabic_option is not None or numeric_option is not None


# =========================
# تحويل الاختيار الرقمي إلى اختيار بحرف
# مثال:
# 1. A1 => أ. A1
# 2. B  => ب. B
# =========================
def normalize_option_line(text):
    text = clean_text(text)

    numeric_match = re.match(r"^([1-4])\s*[\.\-\)]\s*(.*)$", text)

    if numeric_match:
        number = numeric_match.group(1)
        option_text = numeric_match.group(2).strip()
        letter = NUMBER_TO_LETTER[number]
        return f"{letter}. {option_text}"

    return text


# =========================
# استخراج رقم السؤال
# =========================
def extract_question_id(text):
    match = re.match(r"^(\d+\.\d+)", text)
    if match:
        return match.group(1)
    return None


# =========================
# استخراج رقم الفصل
# 5.12 => 5
# =========================
def extract_topic(question_id):
    return question_id.split(".")[0]


# =========================
# البحث عن صورة السؤال
# images/5.12.jpg
# images/5.12.png
# =========================
def find_question_image(question_id):
    possible_extensions = ["jpg", "jpeg", "png", "webp"]

    for ext in possible_extensions:
        image_path = os.path.join(IMAGES_FOLDER, f"{question_id}.{ext}")
        if os.path.exists(image_path):
            return image_path

    return None


# =========================
# قراءة الأسئلة من questions.docx
# =========================
def read_questions(file_path):
    doc = Document(file_path)

    questions = []
    current_question = None

    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)

        if not text:
            continue

        # بداية سؤال جديد
        if is_question_line(text):
            if current_question:
                questions.append(current_question)

            question_id = extract_question_id(text)

            current_question = {
                "id": question_id,
                "topic": extract_topic(question_id),
                "question": text,
                "options": [],
                "image": find_question_image(question_id),
                "answer_letter": None,
                "answer_index": None
            }

        # اختيار
        elif current_question and is_option_line(text):
            option = normalize_option_line(text)
            current_question["options"].append(option)

        # نص تابع للسؤال
        elif current_question:
            current_question["question"] += "\n" + text

    if current_question:
        questions.append(current_question)

    return questions


# =========================
# تحميل قاعدة الأسئلة وربطها بالإجابات
# =========================
def load_questions_database():
    answers = read_answers(ANSWERS_FILE)
    questions = read_questions(QUESTIONS_FILE)

    final_questions = []

    print(f"✅ تم قراءة {len(answers)} إجابة من ملف answers.docx")
    print(f"✅ تم قراءة {len(questions)} سؤال من ملف questions.docx")

    for q in questions:
        q_id = q["id"]

        if q_id not in answers:
            print(f"⚠️ لا توجد إجابة للسؤال: {q_id}")
            continue

        if len(q["options"]) < 4:
            print(f"⚠️ السؤال {q_id} لا يحتوي 4 اختيارات")
            continue

        q["answer_letter"] = answers[q_id]["letter"]
        q["answer_index"] = answers[q_id]["index"]

        final_questions.append(q)

    return final_questions


# =========================
# ترتيب الأسئلة حسب الفصول
# =========================
def group_questions_by_topic(questions):
    topics = {}

    for q in questions:
        topic = q["topic"]

        if topic not in topics:
            topics[topic] = []

        topics[topic].append(q)

    return topics


# =========================
# اختبار الملف
# =========================
if __name__ == "__main__":
    questions_db = load_questions_database()
    topics_db = group_questions_by_topic(questions_db)

    print(f"\n✅ تم تحميل {len(questions_db)} سؤال صالح")

    for topic, questions in sorted(topics_db.items(), key=lambda x: int(x[0])):
        print(f"📘 الفصل {topic}: {len(questions)} سؤال")

    if questions_db:
        print("\nمثال على أول سؤال:")
        print(questions_db[0])