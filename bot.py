import re
import pytesseract
from PIL import Image
import io
from telegram import Update
from telegram.ext import ApplicationBuilder, MessageHandler, filters, ContextTypes
from rapidfuzz import process
from docx import Document

# ====== الإعدادات ======
TOKEN = ""
# تأكد من أن مسار Tesseract صحيح في جهازك
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# ====== تنظيف (كما في كودك الأصلي) ======
def normalize(text):
    if not text: return ""
    text = text.lower()
    text = re.sub(r'[^\w\s]', ' ', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

# ====== قراءة الأسئلة (من ملفك questions.docx) ======
def load_questions():
    try:
        doc = Document("questions.docx")
        full_text = "\n".join([p.text for p in doc.paragraphs])
        # التقاط النمط 5.1 أو 15.1
        pattern = r'(\d+\.\d+)\s+(.*?)(?=\n\d+\.\d+|\Z)'
        matches = re.findall(pattern, full_text, re.DOTALL)
        
        questions = []
        for q_id, q_text in matches:
            questions.append({
                "id": q_id.strip(),
                "text": q_text.strip(),
                "norm": normalize(q_text)
            })
        print(f"📚 تم تحميل {len(questions)} سؤال.")
        return questions
    except Exception as e:
        print(f"❌ خطأ في تحميل الأسئلة: {e}")
        return []

# ====== قراءة الإجابات (من ملفك answers.docx) ======
def load_answers():
    try:
        doc = Document("answers.docx")
        answers = {}
        # القراءة من النصوص والجداول
        for p in doc.paragraphs:
            match = re.match(r'(\d+\.\d+)\s+(.*)', p.text.strip())
            if match: answers[match.group(1)] = match.group(2).strip()
            
        for table in doc.tables:
            for row in table.rows:
                if len(row.cells) >= 2:
                    left, right = row.cells[0].text.strip(), row.cells[1].text.strip()
                    if re.match(r'\d+\.\d+', left): answers[left] = right
                    elif re.match(r'\d+\.\d+', right): answers[right] = left
        print(f"📊 تم تحميل {len(answers)} إجابة.")
        return answers
    except Exception as e:
        print(f"❌ خطأ في تحميل الإجابات: {e}")
        return {}

# ====== معالجة البيانات ======
questions_data = load_questions()
answers_data = load_answers()

final_data = []
for q in questions_data:
    final_data.append({
        "id": q["id"],
        "text": q["text"],
        "norm": q["norm"],
        "answer": answers_data.get(q["id"], "❌ غير موجودة")
    })

# قائمة النصوص للبحث السريع
search_corpus = [q["norm"] for q in final_data]

# ====== وظيفة البحث ======
async def perform_search(update, input_text):
    norm_text = normalize(input_text)
    if len(norm_text) < 5: return # تجاهل النصوص القصيرة جداً
    
    # البحث باستخدام منطق كودك الأصلي
    result = process.extractOne(norm_text, search_corpus)
    if result:
        match_text, score, idx = result
        if score > 75: # رفعنا الدقة قليلاً لضمان الجودة
            q = final_data[idx]
            await update.message.reply_text(
                f"📌 **تم العثور على سؤال مشابه (رقم {q['id']}):**\n\n{q['text']}\n\n✅ **الإجابة:** {q['answer']}",
                parse_mode="Markdown"
            )
            return True
    return False

# ====== التعامل مع الرسائل ======
async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.message
    if not msg: return

    # 1. إذا كانت الرسالة نصية
    if msg.text:
        text = msg.text.strip()
        # إذا كان رقم سؤال مباشر (مثل 5.1)
        if re.match(r'^\d+\.\d+$', text):
            found = False
            for q in final_data:
                if q["id"] == text:
                    await msg.reply_text(f"📌 سؤال رقم {q['id']}:\n{q['text']}\n\n✅ الإجابة: {q['answer']}")
                    found = True
                    break
            if not found: await msg.reply_text("❌ الرقم غير موجود.")
        else:
            if not await perform_search(update, text):
                await msg.reply_text("❌ لم أجد سؤالاً مطابقاً.")

    # 2. إذا كانت الرسالة صورة (جديد)
    elif msg.photo:
        await msg.reply_text("📸 جاري قراءة الصورة والبحث عن السؤال...")
        photo_file = await msg.photo[-1].get_file()
        img_bytes = await photo_file.download_as_bytearray()
        image = Image.open(io.BytesIO(img_bytes))
        
        # قراءة النص من الصورة (عربي + عبري + إنجليزي)
        extracted_text = pytesseract.image_to_string(image, lang='ara+heb+eng')
        
        if not await perform_search(update, extracted_text):
            await msg.reply_text("❌ تعذر العثور على السؤال من الصورة.")

# ====== التشغيل ======
if __name__ == "__main__":
    app = ApplicationBuilder().token(TOKEN).build()
    # الفلتر الآن يقبل النص والصور
    app.add_handler(MessageHandler(filters.TEXT | filters.PHOTO, handle_message))
    print("🤖 البوت يعمل الآن وجاهز لاستقبال النصوص والصور...")
    app.run_polling()
